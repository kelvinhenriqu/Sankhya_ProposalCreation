from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.errors import PdfGenerationError


class SellerDirectory:
    def __init__(self, workbook_path: Path):
        self._workbook_path = workbook_path

    def find(self, seller_name: str) -> dict[str, str]:
        if not self._workbook_path.is_file():
            raise PdfGenerationError(f"Planilha de vendedores não encontrada: {self._workbook_path}")

        workbook = load_workbook(self._workbook_path, read_only=True, data_only=True)
        try:
            expected = seller_name.strip().upper()
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers: dict[str, int] | None = None
                for values in rows:
                    normalized = {
                        str(value).strip().upper(): index
                        for index, value in enumerate(values)
                        if value is not None
                    }
                    if headers is None and "VENDEDOR" in normalized:
                        headers = normalized
                        continue
                    if headers is None:
                        continue
                    seller = self._cell(values, headers.get("VENDEDOR")).strip().upper()
                    if seller == expected:
                        return {
                            "Telefone": self._cell(values, headers.get("TELEFONE")),
                            "Email": self._cell(values, headers.get("EMAIL")),
                        }
            return {"Telefone": "", "Email": ""}
        finally:
            workbook.close()

    @staticmethod
    def _cell(values: tuple[Any, ...], index: int | None) -> str:
        if index is None or index >= len(values) or values[index] is None:
            return ""
        value = values[index]
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

